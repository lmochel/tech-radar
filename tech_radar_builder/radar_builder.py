import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Color
from jinja2 import Environment, FileSystemLoader
import re
import markdown


def markdown_to_html_table(md_file_path):
    """
    Convertit un fichier Markdown en HTML avec une table à une ligne et deux colonnes.
    Utilise '---' comme séparateur entre les colonnes.
    """
    # Lire le contenu du fichier Markdown
    with open(md_file_path, "r", encoding="utf-8") as f:
        md_content = f.read()

    # Séparer en colonnes à partir du séparateur visuel '---'
    # On suppose que chaque '---' correspond à un passage à la colonne suivante
    columns_md = md_content.split('---')

    # Convertir chaque bloc Markdown en HTML
    columns_html = [markdown.markdown(col.strip()) for col in columns_md]

    # Construire la table HTML
    table_html = "<table>\n<tr>\n"
    for col_html in columns_html:
        table_html += f"<td>\n{col_html}\n</td>\n"
    table_html += "</tr>\n</table>"

    return table_html

def clean_field_name(name: str) -> str:
    """
    Nettoie un nom de champ Excel en format snake_case (minuscule, underscore, sans ponctuation).
    Exemple : "Repo url:" -> "repo_url"
    """
    name = str(name).strip().lower()
    name = re.sub(r'[^a-z0-9]+', '_', name)  # remplace tout sauf lettres/chiffres par "_"
    name = re.sub(r'_+', '_', name)           # évite les doublons d'underscore
    return name.strip('_')                    # supprime underscores en trop


def generate_metadata_from_excel(excel_path, sheet_name="Metadata"):
    """
    Lit un onglet Excel contenant deux colonnes : Champ / Valeur
    Exemple :
        | Field        | Value                       |
        |---------------|-----------------------------|
        | Title:        | My Tech Radar               |
        | Description:  | Example radar visualization |
        | Repo url:     | https://github.com/monrepo  |

    Retourne :
        {
          "title": "My Tech Radar",
          "description": "Example radar visualization",
          "repo_url": "https://github.com/monrepo"
        }
    """
    wb = load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name]

    metadata = {}
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row[0] or not row[1]:
            continue
        key = clean_field_name(row[0])
        value = str(row[1]).strip()
        metadata[key] = value

    return metadata


def get_hex_from_fill(cell):
    """
    Extrait la couleur de fond (fill) d'une cellule Excel au format hexadécimal.
    Retourne None si la cellule n'a pas de couleur définie.
    """
    fill = cell.fill
    if fill and fill.fgColor and fill.fgColor.type == "rgb" and fill.fgColor.rgb:
        rgb = fill.fgColor.rgb
        # Supprimer l'alpha s'il existe (ex: 'FF5BA300' -> '#5BA300')
        return f"#{rgb[-6:]}"
    return None


def generate_rings_from_excel(excel_path, sheet_name="Zones"):
    """
    Lit l'onglet 'Zones' et retourne la liste des rings sous la forme :
    [
      { "name": "ADOPT", "color": "#5ba300", "rank": 0 },
      { "name": "TRIAL", "color": "#009eb0", "rank": 1 },
      ...
    ]
    """
    wb = load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name]

    # Trouver les colonnes par nom
    headers = {cell.value: idx for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}

    rings = []
    for row in ws.iter_rows(min_row=2):
        zone_cell = row[headers["Zone"]]
        rank_cell = row[headers["Rank"]]
        color_cell = row[headers["Color"]]

        if zone_cell.value is None or rank_cell.value is None:
            continue

        color_hex = get_hex_from_fill(color_cell) or "#cccccc"  # Valeur par défaut si aucune couleur
        name = str(zone_cell.value).strip()
        rank = int(rank_cell.value)

        rings.append({
            "name": name,
            "color": color_hex,
            "rank": rank
        })

    # Trier les rings par rang croissant
    rings = sorted(rings, key=lambda x: x["rank"])

    return rings

def generate_quadrants_from_excel(excel_path, sheet_name="Sectors"):
    """
    Lit l'onglet 'Sectors' et retourne la liste des quadrants sous forme :
    [
      { "name": "Techniques", "quadrant": 0 },
      { "name": "Outils", "quadrant": 1 },
      ...
    ]
    """
    wb = load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name]

    # Trouver les colonnes par nom
    headers = {cell.value: idx for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}

    quadrants = []
    for row in ws.iter_rows(min_row=2):
        num_cell = row[headers["#"]]
        name_cell = row[headers["Sector"]]

        if num_cell.value is None or name_cell.value is None:
            continue

        # Attention : le radar JS utilise un index de quadrant à partir de 0
        quadrant_index = int(num_cell.value) - 1

        quadrants.append({
            "quadrant": quadrant_index,
            "name": str(name_cell.value).strip()
        })

    # Tri par numéro
    quadrants = sorted(quadrants, key=lambda x: x["quadrant"])

    return quadrants

def generate_entries_from_tech_list(excel_path, sheet_name="Tech. List"):
    """
    Lit l'onglet 'Tech. List' et retourne une liste de dictionnaires avec :
    Active, Technology, Tech. Sector, Monitoring Zone, Trend, Link, Comment.

    Les cellules vides sont converties en None.
    La colonne Active devient True si 'X', sinon False.
    """
    df = pd.read_excel(excel_path, sheet_name=sheet_name)

    # Transformer la colonne Active : 'X' -> True, sinon False
    df["Active"] = df["Active"].apply(lambda x: True if str(x).strip().upper() == "X" else False)

    # Forcer toutes les colonnes en type object pour pouvoir remplacer NaN
    df = df.astype(object)

    # Remplacer tous les NaN par None
    df = df.where(pd.notna(df), None)

    # Convertir en liste de dictionnaires
    entries = df.to_dict(orient="records")

    return entries

def generate_tech_radar_data(excel_path):
    """
    Construit le dictionnaire complet Tech Radar à partir d'un fichier Excel,
    en réutilisant les fonctions déjà existantes :
    - generate_metadata_from_excel
    - generate_quadrants_from_excel
    - generate_rings_from_excel
    - generate_entries_from_tech_list
    """
    # -------------------- METADATA --------------------
    metadata = generate_metadata_from_excel(excel_path)
    
    # -------------------- QUADRANTS --------------------
    quadrants = generate_quadrants_from_excel(excel_path)
    quadrant_to_number = {item['name']: item['quadrant'] for item in quadrants}
    
    # -------------------- RINGS --------------------
    rings = generate_rings_from_excel(excel_path)
    ring_to_number = {item['name']: item['rank'] for item in rings}
    for item in rings:
        if "name" in item and item["name"] is not None:
            item["name"] = item["name"].upper()
    
    # -------------------- ENTRIES --------------------
    entries_raw = generate_entries_from_tech_list(excel_path)
    
    # Mapping des colonnes pour correspondre au format Tech Radar
    trend_to_moved = {
        "Moved out (▼)": -1,
        "No change (●)": 0,
        "Moved in (▲)": 1,
        "New (*)": 2
    }
    
    entries = []
    for row in entries_raw:
        entry = {
            "active": row.get("Active", False),
            "label": row.get("Technology"),
            "quadrant": quadrant_to_number.get((row.get("Tech. Sector"))),
            "ring": ring_to_number.get((row.get("Monitoring Zone"))),
            "link": row.get("Link", None),
            "moved": trend_to_moved.get(str(row.get("Trend")).strip(), 0)
        }
        entries.append(entry)
    
    # -------------------- ASSEMBLE --------------------
    data = {
        "metadata": metadata,
        "quadrants": [{k: v for k, v in q.items() if k != "quadrant"} for q in quadrants],
        "rings": [{k: v for k, v in r.items() if k != "rank"} for r in rings],
        "entries": entries
    }
    
    return data

# ---------- Generate HTML ----------
def generate_index_html(md_file, excel_path, template_path, output_path, css_path=None):
    """
    Génère la page index.html du Tech Radar à partir :
      - du fichier Markdown (description)
      - du fichier Excel (données)
      - du template HTML Jinja2
      - du fichier CSS (optionnel, inline si trouvé)

    Si css_path est fourni et le fichier existe, le CSS est injecté directement dans la page.
    """

    # Charger le template
    template_dir = os.path.dirname(template_path)
    template_file = os.path.basename(template_path)
    env = Environment(loader=FileSystemLoader(template_dir))
    template = env.get_template(template_file)

    # --- Générer les données du radar ---
    data = generate_tech_radar_data(excel_path)

    # --- Convertir le markdown en HTML table ---
    html_table = markdown_to_html_table(md_file)

    # --- Charger le CSS s'il existe ---
    css_content = None
    if css_path and os.path.exists(css_path):
        with open(css_path, "r", encoding="utf-8") as f:
            css_content = f.read()

    # --- Rendu final ---
    rendered_html = template.render(
        metadata=data["metadata"],
        quadrants=data["quadrants"],
        rings=data["rings"],
        entries=data["entries"],
        html_table=html_table,
        css=css_content
    )

    # --- Sauvegarde du fichier ---
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(rendered_html)

    print(f"✅ index.html généré : {output_path}")


class TechRadarBuilder:

    def __init__(self, excel_path, markdown_path):
        self.excel = excel_path
        self.markdown = markdown_path

    